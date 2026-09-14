"""Celery-задача формирования отчёта по партии (Excel/PDF)."""

import asyncio
import os
import tempfile
from datetime import UTC, datetime, timedelta
from typing import Any

from loguru import logger
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from celery_app import celery_app
from core.database import get_session
from data.models.batch import Batch
from domain.services.batch_service import BatchService
from storage.minio_service import MinIOService

_REPORTS_BUCKET = "reports"


@celery_app.task(bind=True, max_retries=3)  # type: ignore[untyped-decorator]
def generate_batch_report(
    self: Any,
    batch_id: int,
    format: str = "excel",
    user_email: str | None = None,
) -> dict[str, Any]:
    """Сформировать отчёт по партии в формате Excel или PDF и загрузить в MinIO."""
    return asyncio.run(_generate_batch_report_async(batch_id, format))


async def _generate_batch_report_async(batch_id: int, format: str) -> dict[str, Any]:
    async with get_session() as session:
        service = BatchService(session)
        batch = await service.get_batch(batch_id)
        stats = _compute_stats(batch)

        ext = "pdf" if format == "pdf" else "xlsx"
        object_name = f"batch_{batch_id}_report.{ext}"

        with tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False) as tmp:
            temp_path = tmp.name

        try:
            if format == "pdf":
                _build_pdf(batch, stats, temp_path)
            else:
                _build_excel(batch, stats, temp_path)

            file_size = os.path.getsize(temp_path)
            file_url = MinIOService().upload_file(
                bucket=_REPORTS_BUCKET,
                file_path=temp_path,
                object_name=object_name,
            )
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    # TODO(webhooks): fire "report_generated" event via WebhookService once wired

    expires_at = (datetime.now(UTC) + timedelta(days=7)).isoformat()

    logger.info(
        "Отчёт по партии batch_id={} сформирован ({}): {}",
        batch_id,
        format,
        object_name,
    )

    return {
        "success": True,
        "file_url": file_url,
        "file_name": object_name,
        "file_size": file_size,
        "expires_at": expires_at,
    }


def _compute_stats(batch: Batch) -> dict[str, Any]:
    """Считает статистику по уже загруженному списку продукции батча."""
    total = len(batch.products)
    aggregated = sum(1 for p in batch.products if p.is_aggregated)
    remaining = total - aggregated
    percent = round(aggregated / total * 100, 2) if total else 0.0
    return {
        "total": total,
        "aggregated": aggregated,
        "remaining": remaining,
        "percent": percent,
    }


def _build_excel(batch: Batch, stats: dict[str, Any], file_path: str) -> None:
    """Строит Excel-отчёт с тремя листами: инфо, продукция, статистика."""
    wb = Workbook()

    info_ws = wb.active
    assert info_ws is not None
    info_ws.title = "Информация о партии"
    _fill_info_sheet(info_ws, batch)

    products_ws = wb.create_sheet("Продукция")
    products_ws.append(["ID", "Уникальный код", "Аггрегирована", "Дата аггрегации"])
    for product in batch.products:
        products_ws.append(
            [
                product.id,
                product.unique_code,
                "Да" if product.is_aggregated else "Нет",
                product.aggregated_at.isoformat() if product.aggregated_at else "",
            ]
        )

    stats_ws = wb.create_sheet("Статистика")
    stats_ws.append(["Всего продукции", stats["total"]])
    stats_ws.append(["Агрегировано", stats["aggregated"]])
    stats_ws.append(["Осталось", stats["remaining"]])
    stats_ws.append(["Процент выполнения", f"{stats['percent']}%"])

    wb.save(file_path)


def _fill_info_sheet(ws: Worksheet, batch: Batch) -> None:
    rows = [
        ("Номер партии", batch.batch_number),
        ("Дата партии", batch.batch_date.isoformat()),
        ("Статус", "Закрыта" if batch.is_closed else "Открыта"),
        ("Рабочий центр", batch.work_center.name),
        ("Смена", batch.shift),
        ("Бригада", batch.team),
        ("Номенклатура", batch.nomenclature),
        ("Начало смены", batch.shift_start.isoformat()),
        ("Окончание смены", batch.shift_end.isoformat()),
    ]
    for label, value in rows:
        ws.append([label, value])


def _build_pdf(batch: Batch, stats: dict[str, Any], file_path: str) -> None:
    """Строит простой табличный PDF-отчёт (без графиков) с теми же тремя секциями."""
    doc = SimpleDocTemplate(file_path, pagesize=A4)
    styles = getSampleStyleSheet()
    elements: list[Any] = []

    elements.append(
        Paragraph(f"Отчёт по партии №{batch.batch_number}", styles["Title"])
    )
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Информация о партии", styles["Heading2"]))
    info_data = [
        ["Номер партии", str(batch.batch_number)],
        ["Дата партии", batch.batch_date.isoformat()],
        ["Статус", "Закрыта" if batch.is_closed else "Открыта"],
        ["Рабочий центр", batch.work_center.name],
        ["Смена", batch.shift],
        ["Бригада", batch.team],
        ["Номенклатура", batch.nomenclature],
        ["Начало смены", batch.shift_start.isoformat()],
        ["Окончание смены", batch.shift_end.isoformat()],
    ]
    elements.append(_styled_table(info_data))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Продукция", styles["Heading2"]))
    products_data = [["ID", "Уникальный код", "Аггрегирована", "Дата аггрегации"]]
    for product in batch.products:
        products_data.append(
            [
                str(product.id),
                product.unique_code,
                "Да" if product.is_aggregated else "Нет",
                product.aggregated_at.isoformat() if product.aggregated_at else "",
            ]
        )
    elements.append(_styled_table(products_data, header=True))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Статистика", styles["Heading2"]))
    stats_data = [
        ["Всего продукции", str(stats["total"])],
        ["Агрегировано", str(stats["aggregated"])],
        ["Осталось", str(stats["remaining"])],
        ["Процент выполнения", f"{stats['percent']}%"],
    ]
    elements.append(_styled_table(stats_data))

    doc.build(elements)


def _styled_table(data: list[list[str]], header: bool = False) -> Table:
    table = Table(data)
    style = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]
    if header:
        style.append(("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey))
        style.append(("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))
    return table
